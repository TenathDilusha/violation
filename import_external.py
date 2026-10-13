"""Convert the externally collected violation articles (xlsx) to data/external_violations.csv.

data/combined_dataset.xlsx has the articles (all English, all violations); text_only_violations.xlsx
adds keyword, clause and excerpt for some of them (matched by URL); data/external_corrections.json
replaces the text of rows the spreadsheet got wrong. None of these articles are
in raw_articles, so each gets a stable id "ext-<uuid5 of its URL, or of its text when there is
no URL>" and status EXTERNAL_VIOLATION. load_to_db.py appends the CSV to violation_articles.

Usage:
    python import_external.py
"""

from __future__ import annotations

import csv
import json
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
ARTICLES_XLSX = ROOT / "data" / "combined_dataset.xlsx"
CLAUSES_XLSX = ROOT / "data" / "text_only_violations.xlsx"
CORRECTIONS = ROOT / "data" / "external_corrections.json"
OUT_CSV = ROOT / "data" / "external_violations.csv"
STATUS = "EXTERNAL_VIOLATION"
COLUMNS = ["article_id", "heading", "content", "lang", "status",
           "source", "url", "outlet", "date_posted", "keyword", "clause", "excerpt"]


def read_sheet(path: Path) -> list[dict]:
    rows = list(load_workbook(path, read_only=True).active.iter_rows(values_only=True))
    header = [str(h).strip() for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def text(value) -> str:
    return "" if value is None else str(value).strip()


def as_date(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
    return text(value)


def article_id(url: str, content: str) -> str:
    key = url if url.startswith("http") else content
    return f"ext-{uuid.uuid5(uuid.NAMESPACE_URL, key)}"


def main() -> None:
    extra = {text(r["url"]): r for r in read_sheet(CLAUSES_XLSX)}
    fixes = json.loads(CORRECTIONS.read_text(encoding="utf-8")) if CORRECTIONS.exists() else {}
    out = []
    for r in read_sheet(ARTICLES_XLSX):
        url, content = text(r["url"]), text(r["content"])
        content = fixes.get(url, {}).get("content", content)
        if not content:
            continue
        x = extra.get(url, {})
        out.append({
            "article_id": article_id(url, content),
            "heading": text(r["heading"]),
            "content": content,
            "lang": text(r["lang"]) or "en",
            "status": STATUS,
            "source": ARTICLES_XLSX.name,
            "url": url,
            "outlet": text(r["outlet"]),
            "date_posted": as_date(r["date_posted"]),
            "keyword": text(x.get("keyword")),
            "clause": text(x.get("clause")),
            "excerpt": text(x.get("excerpt_text")),
        })
    ids = [a["article_id"] for a in out]
    assert len(ids) == len(set(ids)), "duplicate article ids"
    OUT_CSV.parent.mkdir(exist_ok=True)
    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(out)
    print(f"{len(out)} articles ({sum(1 for a in out if a['clause'])} with a clause) written to {OUT_CSV}")


if __name__ == "__main__":
    main()
